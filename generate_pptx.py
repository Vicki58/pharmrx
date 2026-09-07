import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

def create_presentation():
    prs = Presentation()
    # Set slide dimensions to widescreen (16:9)
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    blank_layout = prs.slide_layouts[6]

    # Theme colors
    COLOR_BG = RGBColor(15, 23, 42)       # Slate 900 #0f172a
    COLOR_CARD = RGBColor(30, 41, 59)     # Slate 800 #1e293b
    COLOR_CYAN = RGBColor(6, 182, 212)    # Cyan 500 #06b6d4
    COLOR_WHITE = RGBColor(248, 250, 252) # Slate 50 #f8fafc
    COLOR_MUTED = RGBColor(148, 163, 184) # Slate 400 #94a3b8
    COLOR_BORDER = RGBColor(51, 65, 85)   # Slate 700 #334155
    COLOR_ACCENT = RGBColor(99, 102, 241) # Indigo 500 #6366f1

    def add_slide_background(slide):
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
        bg.fill.solid()
        bg.fill.fore_color.rgb = COLOR_BG
        bg.line.color.rgb = COLOR_BG
        return bg

    def add_header(slide, section_tag, title_text):
        # Section tag (pill)
        tag_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11.7), Inches(0.4))
        tf = tag_box.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = section_tag.upper()
        p.font.size = Pt(11)
        p.font.bold = True
        p.font.color.rgb = COLOR_CYAN

        # Title
        title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.7), Inches(11.7), Inches(0.8))
        tf = title_box.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = title_text
        p.font.size = Pt(26)
        p.font.bold = True
        p.font.color.rgb = COLOR_WHITE

    # ----------------------------------------------------
    # SLIDE 1: Title Slide
    # ----------------------------------------------------
    s1 = prs.slides.add_slide(blank_layout)
    add_slide_background(s1)

    # Decorative card
    card = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.2), Inches(1.2), Inches(10.9), Inches(5.1))
    card.fill.solid()
    card.fill.fore_color.rgb = COLOR_CARD
    card.line.color.rgb = COLOR_BORDER

    # Badge inside card
    badge = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(2.0), Inches(1.8), Inches(4.5), Inches(0.45))
    badge.fill.solid()
    badge.fill.fore_color.rgb = RGBColor(16, 50, 75)
    badge.line.color.rgb = COLOR_CYAN
    p = badge.text_frame.paragraphs[0]
    p.text = "IBL12307: WEB DEVELOPMENT LABORATORY PRACTICAL"
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = COLOR_CYAN
    p.alignment = PP_ALIGN.CENTER

    # Main Title
    t_box = s1.shapes.add_textbox(Inches(2.0), Inches(2.5), Inches(9.3), Inches(1.5))
    tf = t_box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "PharmRx: Pharmacy Management System"
    p.font.size = Pt(36)
    p.font.bold = True
    p.font.color.rgb = COLOR_WHITE

    p2 = tf.add_paragraph()
    p2.text = "A Secure, Role-Based PHP & MySQL Web Application"
    p2.font.size = Pt(18)
    p2.font.color.rgb = COLOR_MUTED

    # Presenter Details
    d_box = s1.shapes.add_textbox(Inches(2.0), Inches(4.3), Inches(9.3), Inches(1.5))
    tf_d = d_box.text_frame
    p_d1 = tf_d.paragraphs[0]
    p_d1.text = "Instructor: Dr. Edwin Ngwawe"
    p_d1.font.size = Pt(14)
    p_d1.font.bold = True
    p_d1.font.color.rgb = COLOR_WHITE

    p_d2 = tf_d.add_paragraph()
    p_d2.text = "Deliverables: Source Code | MySQL Database (.sql) | Report | 10-Min Live Demonstration"
    p_d2.font.size = Pt(13)
    p_d2.font.color.rgb = COLOR_CYAN

    # ----------------------------------------------------
    # SLIDE 2: Introduction (0:00 - 0:45)
    # ----------------------------------------------------
    s2 = prs.slides.add_slide(blank_layout)
    add_slide_background(s2)
    add_header(s2, "0:00 - 0:45 | Section 1", "1. Introduction & Real-World Problem")

    # 2 Cards: Left = Problem, Right = Solution & Target Users
    c1 = s2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.7), Inches(5.6), Inches(5.1))
    c1.fill.solid()
    c1.fill.fore_color.rgb = COLOR_CARD
    c1.line.color.rgb = COLOR_BORDER

    tf1 = c1.text_frame
    tf1.word_wrap = True
    p = tf1.paragraphs[0]
    p.text = "THE REAL-WORLD PROBLEM"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = RGBColor(248, 113, 113) # Red/Salmon

    items1 = [
        "Traditional pharmacies rely on manual paper logbooks and fragmented spreadsheets.",
        "High risk of stock discrepancies, expired medications, and overselling.",
        "Absence of user accountability: no verifiable record of who added, edited, or deleted items.",
        "Security risks: vulnerability to price tampering and unauthorized access."
    ]
    for it in items1:
        p = tf1.add_paragraph()
        p.text = "• " + it
        p.font.size = Pt(13)
        p.font.color.rgb = COLOR_WHITE
        p.space_before = Pt(12)

    c2 = s2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.9), Inches(1.7), Inches(5.6), Inches(5.1))
    c2.fill.solid()
    c2.fill.fore_color.rgb = COLOR_CARD
    c2.line.color.rgb = COLOR_BORDER

    tf2 = c2.text_frame
    tf2.word_wrap = True
    p = tf2.paragraphs[0]
    p.text = "THE PHARMRX SOLUTION & TARGET USERS"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = COLOR_CYAN

    items2 = [
        "Centralized Inventory Hub: Real-time stock levels, pricing in KES, and visual packaging photos.",
        "Automated Sales POS: Instant price calculation and automatic stock deduction.",
        "Target User Role 1: Pharmacy Administrator (Full CRUD control, system audit logs, restock powers).",
        "Target User Role 2: Normal Staff / Pharmacist (View catalog, check stock, record sales transactions)."
    ]
    for it in items2:
        p = tf2.add_paragraph()
        p.text = "✔ " + it
        p.font.size = Pt(13)
        p.font.color.rgb = COLOR_WHITE
        p.space_before = Pt(12)

    # ----------------------------------------------------
    # SLIDE 3: Objectives & Scope (0:45 - 1:30)
    # ----------------------------------------------------
    s3 = prs.slides.add_slide(blank_layout)
    add_slide_background(s3)
    add_header(s3, "0:45 - 1:30 | Section 2", "2. Project Objectives & System Scope")

    # 3 Cards for Objectives
    objs = [
        ("Objective 1: Automated Inventory & Sales", "Replace error-prone manual logs with transactional medicine cataloging, categorized inventory, and automated inventory deduction upon sale."),
        ("Objective 2: Robust Role-Based Access", "Enforce strict privilege separation between Administrators (CRUD + auditing) and Normal Staff (view-only catalog + transaction recording)."),
        ("Objective 3: Security & Full Auditability", "Demonstrate industry best practices: 100% PDO prepared statements, Bcrypt password hashing, XSS sanitization, and automated activity logging.")
    ]
    for i, (title, desc) in enumerate(objs):
        x = Inches(0.8 + i * 4.0)
        card = s3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, Inches(1.7), Inches(3.7), Inches(2.5))
        card.fill.solid()
        card.fill.fore_color.rgb = COLOR_CARD
        card.line.color.rgb = COLOR_BORDER
        tf = card.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = title
        p.font.size = Pt(14)
        p.font.bold = True
        p.font.color.rgb = COLOR_CYAN
        p2 = tf.add_paragraph()
        p2.text = desc
        p2.font.size = Pt(12)
        p2.font.color.rgb = COLOR_WHITE
        p2.space_before = Pt(8)

    # Bottom Scope Card
    b_card = s3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(4.5), Inches(11.7), Inches(2.3))
    b_card.fill.solid()
    b_card.fill.fore_color.rgb = COLOR_CARD
    b_card.line.color.rgb = COLOR_BORDER
    tf_b = b_card.text_frame
    tf_b.word_wrap = True
    p = tf_b.paragraphs[0]
    p.text = "SYSTEM SCOPE SUMMARY"
    p.font.size = Pt(15)
    p.font.bold = True
    p.font.color.rgb = COLOR_ACCENT

    scope_points = [
        "User Auth: Secure Register, Login, Logout, and Change Password.",
        "Catalog Management: Create, Read, Update, Delete medicines with KES currency & image uploads.",
        "POS & Inventory: Multi-item sales records, live price calculator, automatic stock deduction.",
        "Audit Trail: System-wide activity logs tracking User, Action (Login, Logout, Add, Edit, Delete), and Time."
    ]
    for pt in scope_points:
        p = tf_b.add_paragraph()
        p.text = "• " + pt
        p.font.size = Pt(12.5)
        p.font.color.rgb = COLOR_WHITE
        p.space_before = Pt(4)

    # ----------------------------------------------------
    # SLIDE 4: Database Design (1:30 - 2:30)
    # ----------------------------------------------------
    s4 = prs.slides.add_slide(blank_layout)
    add_slide_background(s4)
    add_header(s4, "1:30 - 2:30 | Section 3", "3. Relational Database Design & Normalization (3NF)")

    # 5 Table Cards in a Grid
    tables = [
        ("users", "Authentication & Roles", ["id (PK, Auto)", "username (Unique)", "email (Unique)", "password_hash (Bcrypt)", "role (Admin / Normal)"]),
        ("categories", "Medicine Grouping", ["id (PK, Auto)", "name (Unique)", "description", "created_at"]),
        ("medicines", "Main Inventory Entity", ["id (PK, Auto)", "category_id (FK -> categories.id)", "name (Unique)", "price (KES), stock_quantity", "image_path"]),
        ("sales", "Transaction Bridge Entity", ["id (PK, Auto)", "medicine_id (FK -> medicines.id)", "user_id (FK -> users.id)", "quantity, total_price (KES)", "sale_date"]),
        ("activity_logs", "Security Audit Trail", ["id (PK, Auto)", "user_id (FK -> users.id)", "activity (ENUM)", "description (Action details)", "log_time"])
    ]

    for i, (tname, tdesc, fields) in enumerate(tables):
        if i < 3:
            x = Inches(0.8 + i * 4.0)
            y = Inches(1.7)
            w = Inches(3.7)
            h = Inches(2.5)
        else:
            x = Inches(0.8 + (i - 3) * 6.0)
            y = Inches(4.4)
            w = Inches(5.7)
            h = Inches(2.4)

        card = s4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h)
        card.fill.solid()
        card.fill.fore_color.rgb = COLOR_CARD
        card.line.color.rgb = COLOR_BORDER
        tf = card.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = f"TABLE: `{tname}`"
        p.font.size = Pt(13)
        p.font.bold = True
        p.font.color.rgb = COLOR_CYAN

        p_sub = tf.add_paragraph()
        p_sub.text = tdesc
        p_sub.font.size = Pt(11)
        p_sub.font.color.rgb = COLOR_MUTED

        for f in fields:
            p_f = tf.add_paragraph()
            p_f.text = "• " + f
            p_f.font.size = Pt(11)
            p_f.font.color.rgb = COLOR_WHITE

    # ----------------------------------------------------
    # SLIDE 5: System Architecture (2:30 - 3:15)
    # ----------------------------------------------------
    s5 = prs.slides.add_slide(blank_layout)
    add_slide_background(s5)
    add_header(s5, "2:30 - 3:15 | Section 4", "4. System Architecture & Technology Stack")

    techs = [
        ("Backend: PHP 8.3 (PDO)", "Clean modular PHP architecture using PDO connection handles, prepared statement queries, and session management guards."),
        ("Database: MySQL (InnoDB)", "Full relational engine supporting Foreign Key cascading, constraints, and atomic ACID transactions during sales."),
        ("Frontend & UI", "Semantic HTML5, Vanilla JavaScript (client-side validation & live price calculator), and custom dark-theme responsive CSS."),
        ("Security Stack", "Bcrypt hashing (`password_hash`), XSS neutralization (`htmlspecialchars`), HTTP-only secure session cookies.")
    ]

    for i, (title, desc) in enumerate(techs):
        x = Inches(0.8 + (i % 2) * 6.0)
        y = Inches(1.7 + (i // 2) * 2.6)
        card = s5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, Inches(5.7), Inches(2.3))
        card.fill.solid()
        card.fill.fore_color.rgb = COLOR_CARD
        card.line.color.rgb = COLOR_BORDER
        tf = card.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = title
        p.font.size = Pt(15)
        p.font.bold = True
        p.font.color.rgb = COLOR_CYAN
        p2 = tf.add_paragraph()
        p2.text = desc
        p2.font.size = Pt(13)
        p2.font.color.rgb = COLOR_WHITE
        p2.space_before = Pt(8)

    # ----------------------------------------------------
    # SLIDE 6: Live Demonstration Checklist (3:15 - 7:30)
    # ----------------------------------------------------
    s6 = prs.slides.add_slide(blank_layout)
    add_slide_background(s6)
    add_header(s6, "3:15 - 7:30 | Section 5", "5. Live Demonstration: Step-by-Step Flow")

    # 2 Columns of Demo Steps
    demo_steps_left = [
        ("Step 1: Admin Login", "Log in as `admin` (pass: `admin123`). Show KPI metrics (Total Medicines, Revenue in KES, Stock warnings)."),
        ("Step 2: Role Restriction", "Log in as `staff` (pass: `staff123`). Show view-only medicine catalog; direct access to `medicine_add.php` shows 403 Forbidden."),
        ("Step 3: Main CRUD Cycle", "Add new medicine -> View in grid -> Edit details -> Delete record (cleaned up from DB & storage)."),
        ("Step 4: Search & Pagination", "Search keyword 'Amoxicillin' -> Instant filtered cards -> Clear search -> Paginate through pages.")
    ]

    demo_steps_right = [
        ("Step 5: File Upload & Validation", "Upload packaging image (JPG/PNG/WEBP). Demonstrate file size limit (>2MB) rejection."),
        ("Step 6: Form Validation", "Submit blank fields or negative price/stock -> Show client-side JS alert & server-side PHP feedback."),
        ("Step 7: Sales & Stock Deduction", "Record a sale as Staff -> Live price calculator -> Confirm stock auto-decrements in database."),
        ("Step 8: Audit Activity Log", "Open System Activity Logs -> Verify entries for Login, Logout, Add, Edit, Delete with timestamps.")
    ]

    for i, (title, desc) in enumerate(demo_steps_left):
        y = Inches(1.7 + i * 1.3)
        c = s6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), y, Inches(5.6), Inches(1.15))
        c.fill.solid()
        c.fill.fore_color.rgb = COLOR_CARD
        c.line.color.rgb = COLOR_BORDER
        tf = c.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = title
        p.font.size = Pt(13)
        p.font.bold = True
        p.font.color.rgb = COLOR_CYAN
        p2 = tf.add_paragraph()
        p2.text = desc
        p2.font.size = Pt(11)
        p2.font.color.rgb = COLOR_WHITE

    for i, (title, desc) in enumerate(demo_steps_right):
        y = Inches(1.7 + i * 1.3)
        c = s6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.9), y, Inches(5.6), Inches(1.15))
        c.fill.solid()
        c.fill.fore_color.rgb = COLOR_CARD
        c.line.color.rgb = COLOR_BORDER
        tf = c.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = title
        p.font.size = Pt(13)
        p.font.bold = True
        p.font.color.rgb = COLOR_ACCENT
        p2 = tf.add_paragraph()
        p2.text = desc
        p2.font.size = Pt(11)
        p2.font.color.rgb = COLOR_WHITE

    # ----------------------------------------------------
    # SLIDE 7: Security & Testing Evidence (7:30 - 8:30)
    # ----------------------------------------------------
    s7 = prs.slides.add_slide(blank_layout)
    add_slide_background(s7)
    add_header(s7, "7:30 - 8:30 | Section 6", "6. Security Architecture & Test Evidence")

    c_sec = s7.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.7), Inches(5.6), Inches(5.1))
    c_sec.fill.solid()
    c_sec.fill.fore_color.rgb = COLOR_CARD
    c_sec.line.color.rgb = COLOR_BORDER
    tf_sec = c_sec.text_frame
    tf_sec.word_wrap = True
    p = tf_sec.paragraphs[0]
    p.text = "CORE SECURITY CONTROLS"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = COLOR_CYAN

    sec_items = [
        "SQL Injection Protection: Parameterized queries across all 100% of database interactions via PDO.",
        "Password Security: Cryptographic one-way Bcrypt hashing with auto-salting (`PASSWORD_DEFAULT`).",
        "XSS Neutralization: Mandatory `htmlspecialchars()` escaping on every dynamic string output.",
        "Session Hardening: `session_regenerate_id()` on login to prevent Session Fixation, plus HTTP-only flags."
    ]
    for s_it in sec_items:
        p = tf_sec.add_paragraph()
        p.text = "🛡️ " + s_it
        p.font.size = Pt(13)
        p.font.color.rgb = COLOR_WHITE
        p.space_before = Pt(12)

    c_test = s7.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.9), Inches(1.7), Inches(5.6), Inches(5.1))
    c_test.fill.solid()
    c_test.fill.fore_color.rgb = COLOR_CARD
    c_test.line.color.rgb = COLOR_BORDER
    tf_test = c_test.text_frame
    tf_test.word_wrap = True
    p = tf_test.paragraphs[0]
    p.text = "AUTOMATED TEST SUITE (`test_runner.php`)"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = RGBColor(52, 211, 153) # Green

    test_results = [
        "Database PDO Connection ............................. [ PASS ]",
        "Verify Required Table Entities Count (Min 4) ....... [ PASS ]",
        "Password Hashing Integration (Bcrypt) .............. [ PASS ]",
        "User Role Check Helpers (Admin vs Normal) .......... [ PASS ]",
        "Structural Normalization (Foreign Keys) ............ [ PASS ]",
        "Activity Logger Writes to DB ....................... [ PASS ]",
        "Stock & Price Constraints Logic .................... [ PASS ]",
        "",
        "SUMMARY: 7 PASSED, 0 FAILED (100% SUCCESS)"
    ]
    for tr in test_results:
        p = tf_test.add_paragraph()
        p.text = tr
        p.font.size = Pt(12)
        p.font.color.rgb = RGBColor(52, 211, 153) if "[ PASS ]" in tr or "100%" in tr else COLOR_WHITE

    # ----------------------------------------------------
    # SLIDE 8: Challenges & Solutions (8:30 - 9:15)
    # ----------------------------------------------------
    s8 = prs.slides.add_slide(blank_layout)
    add_slide_background(s8)
    add_header(s8, "8:30 - 9:15 | Section 7", "7. Genuine Engineering Challenges & Solutions")

    challenges = [
        ("Challenge 1: Transactional Stock Deductions",
         "PROBLEM: Concurrent sales could cause race conditions or negative inventory if two users purchased simultaneously.\n"
         "SOLUTION: Implemented MySQL InnoDB transactions (`beginTransaction()`, `FOR UPDATE` row lock, `commit()`, and `rollBack()`)."),
        ("Challenge 2: Orphaned Media Files on Disk",
         "PROBLEM: Editing medicine images or deleting records left unused image files consuming server storage.\n"
         "SOLUTION: Programmed an automatic cleanup routine using PHP's `unlink()` to permanently remove old files whenever updated or deleted."),
        ("Challenge 3: Multi-Tier Validation & Security",
         "PROBLEM: Relying solely on client-side JS validation is unsafe, but server-only validation causes poor user experience.\n"
         "SOLUTION: Built synchronized dual-tier validation: JavaScript provides immediate visual feedback, while PHP enforces strict server-side validation.")
    ]

    for i, (title, text) in enumerate(challenges):
        y = Inches(1.7 + i * 1.7)
        c = s8.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), y, Inches(11.7), Inches(1.5))
        c.fill.solid()
        c.fill.fore_color.rgb = COLOR_CARD
        c.line.color.rgb = COLOR_BORDER
        tf = c.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = title
        p.font.size = Pt(14)
        p.font.bold = True
        p.font.color.rgb = COLOR_CYAN
        p2 = tf.add_paragraph()
        p2.text = text
        p2.font.size = Pt(12)
        p2.font.color.rgb = COLOR_WHITE
        p2.space_before = Pt(4)

    # ----------------------------------------------------
    # SLIDE 9: Lessons Learned & Conclusion (9:15 - 10:00)
    # ----------------------------------------------------
    s9 = prs.slides.add_slide(blank_layout)
    add_slide_background(s9)
    add_header(s9, "9:15 - 10:00 | Section 8", "8. Lessons Learned & Future Roadmap")

    c_ll = s9.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.7), Inches(5.6), Inches(5.1))
    c_ll.fill.solid()
    c_ll.fill.fore_color.rgb = COLOR_CARD
    c_ll.line.color.rgb = COLOR_BORDER
    tf_ll = c_ll.text_frame
    tf_ll.word_wrap = True
    p = tf_ll.paragraphs[0]
    p.text = "KEY LESSONS LEARNED"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = COLOR_CYAN

    lessons = [
        "Database normalization (3NF) upfront eliminates costly refactoring and data anomalies.",
        "Parameterized queries and prepared statements should be a fundamental reflex, never an afterthought.",
        "Activity logging provides invaluable forensic insight into multi-user systems.",
        "Role-based routing must be enforced at both the UI presentation layer and the server-side controller."
    ]
    for l_it in lessons:
        p = tf_ll.add_paragraph()
        p.text = "💡 " + l_it
        p.font.size = Pt(13)
        p.font.color.rgb = COLOR_WHITE
        p.space_before = Pt(12)

    c_fr = s9.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.9), Inches(1.7), Inches(5.6), Inches(5.1))
    c_fr.fill.solid()
    c_fr.fill.fore_color.rgb = COLOR_CARD
    c_fr.line.color.rgb = COLOR_BORDER
    tf_fr = c_fr.text_frame
    tf_fr.word_wrap = True
    p = tf_fr.paragraphs[0]
    p.text = "FUTURE IMPROVEMENTS"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = COLOR_ACCENT

    future = [
        "Barcode Scanning: Rapid point-of-sale checkout using webcam barcode scanner API.",
        "Automated Notifications: SMS/Email alerts to suppliers when stock levels drop below reorder thresholds.",
        "Receipt Printing: Direct PDF invoice & receipt generation for customer purchases.",
        "Supplier Module: Full purchase order tracking and supplier delivery ledger."
    ]
    for f_it in future:
        p = tf_fr.add_paragraph()
        p.text = "🚀 " + f_it
        p.font.size = Pt(13)
        p.font.color.rgb = COLOR_WHITE
        p.space_before = Pt(12)

    # ----------------------------------------------------
    # SLIDE 10: Conclusion & Q&A
    # ----------------------------------------------------
    s10 = prs.slides.add_slide(blank_layout)
    add_slide_background(s10)

    card10 = s10.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.5), Inches(1.5), Inches(10.3), Inches(4.5))
    card10.fill.solid()
    card10.fill.fore_color.rgb = COLOR_CARD
    card10.line.color.rgb = COLOR_CYAN

    tf10 = card10.text_frame
    tf10.word_wrap = True
    p = tf10.paragraphs[0]
    p.text = "THANK YOU!"
    p.font.size = Pt(36)
    p.font.bold = True
    p.font.color.rgb = COLOR_CYAN
    p.alignment = PP_ALIGN.CENTER

    p2 = tf10.add_paragraph()
    p2.text = "PharmRx Pharmacy Management System"
    p2.font.size = Pt(22)
    p2.font.bold = True
    p2.font.color.rgb = COLOR_WHITE
    p2.alignment = PP_ALIGN.CENTER
    p2.space_before = Pt(10)

    p3 = tf10.add_paragraph()
    p3.text = "100% Rubric Coverage: Auth (10) | Roles (10) | Database (10) | CRUD (15) | Search/Pagination (10) | Upload (10) | Validation (10) | Security (10) | Logs (5) | UI (5) | Testing (5)"
    p3.font.size = Pt(13)
    p3.font.color.rgb = COLOR_MUTED
    p3.alignment = PP_ALIGN.CENTER
    p3.space_before = Pt(16)

    p4 = tf10.add_paragraph()
    p4.text = "Ready for Questions & Live Code Demonstration"
    p4.font.size = Pt(18)
    p4.font.bold = True
    p4.font.color.rgb = COLOR_WHITE
    p4.alignment = PP_ALIGN.CENTER
    p4.space_before = Pt(20)

    out_path = r"C:\Users\incai\Web_dev_practical\Web_Final project\Pharmacy_Management_System_Presentation.pptx"
    prs.save(out_path)
    print("SUCCESS: Presentation saved to", out_path)

if __name__ == "__main__":
    create_presentation()
